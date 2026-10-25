"""将 docs/使用手册.md 转换为带图片的 Word 文档。

用法：python scripts/md-to-docx.py
输出：docs/XHCMS小说站使用手册.docx
"""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "docs" / "使用手册.md"
OUT = ROOT / "docs" / "XHCMS小说站使用手册.docx"

BODY_FONT = "微软雅黑"
CODE_FONT = "Consolas"
ACCENT = RGBColor(0x63, 0x66, 0xF1)
MUTED = RGBColor(0x6B, 0x72, 0x80)


def set_east_asian(run_or_style, font_name: str) -> None:
    font = run_or_style.font
    font.name = font_name
    rpr = run_or_style.element.get_or_add_rPr() if hasattr(run_or_style, "element") else run_or_style._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)


def shade(cell_or_paragraph, color: str) -> None:
    element = cell_or_paragraph._tc if hasattr(cell_or_paragraph, "_tc") else cell_or_paragraph._p
    pr = element.get_or_add_tcPr() if hasattr(cell_or_paragraph, "_tc") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    pr.append(shd)


def add_runs(paragraph, text: str, base_size: int = 10.5, bold: bool = False) -> None:
    """处理行内 **粗体**、`代码`、[文本](链接)。"""
    pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
    for chunk in pattern.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
            run.font.size = Pt(base_size)
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt(base_size - 0.5)
            run.font.color.rgb = RGBColor(0x9A, 0x34, 0x12)
        elif chunk.startswith("[") and "](" in chunk:
            label, target = re.match(r"\[([^\]]+)\]\(([^)]+)\)", chunk).groups()  # type: ignore[union-attr]
            if target.startswith("#"):
                text_value = label
            else:
                text_value = f"{label}（{target}）"
            run = paragraph.add_run(text_value)
            run.font.size = Pt(base_size)
            run.font.color.rgb = ACCENT
        else:
            run = paragraph.add_run(chunk)
            run.font.size = Pt(base_size)
            run.bold = bold


def table_from(block: list[str], doc: Document) -> None:
    rows = []
    for line in block:
        stripped = line.strip().strip("|")
        cells = [cell.strip() for cell in stripped.split("|")]
        rows.append(cells)
    if len(rows) < 2:
        return
    header, body = rows[0], rows[2:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, text in enumerate(header):
        cell = table.rows[0].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        add_runs(paragraph, text, base_size=9.5, bold=True)
    for row in body:
        cells = table.add_row().cells
        for index, raw in enumerate(row[: len(header)]):
            cell = cells[index]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            image = re.search(r"!\[[^\]]*\]\(([^)]+)\)", raw)
            if image:
                image_path = (MD.parent / image.group(1)).resolve()
                if image_path.exists():
                    run = paragraph.add_run()
                    run.add_picture(str(image_path), width=Cm(7.2))
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                add_runs(paragraph, raw, base_size=9.5)
    doc.add_paragraph()


def code_block(block: list[str], doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.4)
    paragraph.paragraph_format.space_after = Pt(4)
    shade(paragraph, "F4F5F7")
    for index, line in enumerate(block):
        run = paragraph.add_run(line)
        run.font.name = CODE_FONT
        run.font.size = Pt(9)
        if index < len(block) - 1:
            run.add_break()


def main() -> None:
    text = MD.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    set_east_asian(normal, BODY_FONT)

    for style_name, size, color in (("Heading 1", 20, ACCENT), ("Heading 2", 15, ACCENT), ("Heading 3", 12.5, RGBColor(0x1F, 0x29, 0x37))):
        style = doc.styles[style_name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.color.rgb = color
        set_east_asian(style, BODY_FONT)
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(6)

    # 封面
    cover_title = doc.add_paragraph()
    cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cover_title.add_run("XHCMS Novel 2.0")
    run.font.size = Pt(30)
    run.bold = True
    run.font.color.rgb = ACCENT
    cover_sub = doc.add_paragraph()
    cover_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cover_sub.add_run("小说网站 · 使用手册")
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    for line in (
        "NestJS 11 + Next.js 15 + TypeScript + Prisma",
        "版本 2.0.0 ｜ 2026-10",
        "",
        "本手册包含：快速开始 / 读者门户 / 作者中心 / 管理后台 / 采集中心 / 部署运维 / FAQ",
    ):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(line)
        run.font.size = Pt(11)
        run.font.color.rgb = MUTED

    cover_image = ROOT / "screenshots" / "01-home.png"
    if cover_image.exists():
        doc.add_paragraph()
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(cover_image), width=Cm(16))
    doc.add_page_break()

    index = 0
    in_code = False
    code_lines: list[str] = []
    table_lines: list[str] = []
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_lines
        if not paragraph_lines:
            return
        merged = " ".join(item.strip() for item in paragraph_lines).strip()
        paragraph_lines = []
        if not merged:
            return
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.25
        add_runs(paragraph, merged)

    def flush_table() -> None:
        nonlocal table_lines
        if table_lines:
            table_from(table_lines, doc)
            table_lines = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                code_block(code_lines, doc)
                code_lines, in_code = [], False
            else:
                flush_paragraph()
                flush_table()
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            table_lines.append(line)
            index += 1
            continue
        if table_lines:
            flush_table()

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        if stripped == "---":
            flush_paragraph()
            index += 1
            continue

        if stripped.startswith("#"):
            flush_paragraph()
            level = len(stripped) - len(stripped.lstrip("#"))
            title = stripped[level:].strip()
            if title == "XHCMS Novel 2.0 使用手册":
                index += 1
                continue
            doc.add_heading(title, level=min(level, 3))
            index += 1
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            quote = stripped.lstrip("> ").strip()
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.5)
            shade(paragraph, "EEF2FF")
            add_runs(paragraph, quote, base_size=10)
            index += 1
            continue

        image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image:
            flush_paragraph()
            image_path = (MD.parent / image.group(2)).resolve()
            if image_path.exists():
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(image_path), width=Cm(15.5))
                caption = image.group(1)
                if caption:
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = cap.add_run(caption)
                    run.font.size = Pt(9)
                    run.font.color.rgb = MUTED
            index += 1
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if bullet or numbered:
            flush_paragraph()
            content = bullet.group(1) if bullet else numbered.group(1)  # type: ignore[union-attr]
            style = "List Bullet" if bullet else "List Number"
            paragraph = doc.add_paragraph(style=style)
            add_runs(paragraph, content)
            index += 1
            continue

        paragraph_lines.append(line)
        index += 1

    flush_paragraph()
    flush_table()

    property_section = doc.core_properties
    property_section.title = "XHCMS Novel 2.0 使用手册"
    property_section.author = "XHCMS Team"
    property_section.subject = "小说网站使用与运维手册"
    property_section.created = dt.datetime(2026, 10, 25)
    property_section.keywords = "XHCMS,小说站,NestJS,Next.js,使用手册"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    size_kb = OUT.stat().st_size / 1024
    print(f"generated: {OUT} ({size_kb:.0f} KB)")


if __name__ == "__main__":
    main()
